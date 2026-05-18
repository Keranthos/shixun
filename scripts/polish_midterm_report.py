# -*- coding: utf-8 -*-
"""润色中期报告：引号、措辞，并同步更新 Word。"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "实训中期报告.md"
DOCX_PATH = ROOT / "实训中期报告.docx"


def polish_text(text: str) -> str:
    # 1. 直角引号 → 中文双引号
    text = text.replace("「", "“").replace("」", "”")

    # 2. 删除/改写“像 AI 写给自己看的”元说明
    text = re.sub(
        r"\*\*使用方式\*\*：[^\n]+\n",
        "",
        text,
    )
    text = text.replace("（超详版）", "")
    text = text.replace(
        "（含接口表、规划对照、A01–A22 说明书与模块走查），依据《软工资源平台》与当前仓库整理。",
        "依据《软工资源平台》项目说明与当前代码仓库整理。",
    )

    replacements = [
        ("可答辩、可截图、可叙述创新点的完整形态", "功能完整、便于展示与技术总结的完整形态"),
        ("服务综合实训答辩与课程报告撰写", "支撑综合实训中期检查与后续课程报告撰写"),
        ("| Docker Compose | 答辩环境一键复现，降低评审机配置成本 |", "| Docker Compose | 容器化一键部署，降低环境配置成本 |"),
        ("适合答辩演示第一印象", "提升首屏视觉与交互体验"),
        ("## 六、中期后重点完成工作（拓展阶段 · 答辩重点）", "## 六、中期后重点完成工作（拓展阶段）"),
        ("> **建议答辩时**：用 60% 时间展示本章，40% 时间演示第五章主线。\n", ""),
        ("适合答辩“知识关联”话题", "便于呈现资源之间的关联关系"),
        ("保证答辩现场无 Key 也能演示完整交互", "在未配置 API Key 时仍可完成完整交互演示"),
        ("## 七、典型业务场景说明（便于答辩叙述）", "## 七、典型业务场景说明"),
        ("使系统不仅“能用”，而且“好演示、好讲解、好写报告”", "使系统在可用性、可维护性与可扩展性上达到预期目标"),
        ("为期末验收与答辩做好铺垫", "为期末验收与后续功能迭代做好准备"),
        ("答辩 PPT 建议按编号制作“功能清单”一页", "亦可按功能编号对照《高级功能与报告亮点清单》整理附录"),
        ("## 十四、后端 REST API 接口清单（主干）", "## 十四、后端 REST API 接口清单"),
        ("## 十七、高级功能 A01–A22 实现说明书（答辩用）", "## 十七、高级功能 A01–A22 实现说明书"),
        ("以下每项按 **“功能目标 → 用户可见行为 → 关键文件 → 报告可写创新点”** 展开", "以下每项按“功能目标 → 用户可见行为 → 关键文件 → 创新点说明”展开"),
        ("一键部署答辩环境", "一键部署完整运行环境"),
        ("## 十九、答辩现场演示脚本建议（约 8–10 分钟）", "## 十九、系统演示流程说明"),
        ("## 二十四、答辩常见问题与参考回答（FAQ）", "## 二十四、常见问题说明"),
        ("保证答辩可复现", "保证在无密钥环境下仍可复现"),
        ("便于中期检查老师对照“立项书—实现”一致性", "便于对照立项书与实现的一致性"),
        ("报告可写“可替换为 LLM 摘要 API”", "后续可替换为 LLM 摘要 API"),
        ("形成“双端一致”的安全叙事", "实现前后端一致的内容安全校验"),
        ("强化“准实时运营”演示效果", "体现准实时运营反馈效果"),
        ("*（全文完 · 超详版：第一至十二章正文 + 第十三至二十一章附录 + 第二十二至二十六章扩展走查与 FAQ）*", "*（全文完）*"),
        ("演示彩排", "集成验证"),
    ]

    for old, new in replacements:
        text = text.replace(old, new)

    # 鉴权说明中的引号已在上面统一为 “”
    text = text.replace("标注“需登录”", '标注“需登录”')  # no-op, keep

    # 第三章标题里“便于报告中说明可演进性”——去掉 meta
    text = text.replace(
        "便于报告中说明“可演进性”",
        "便于后续扩展为更完整的账号与合规能力",
    )

    return text


def patch_docx(text_rules: list[tuple[str, str]]) -> None:
    from docx import Document

    doc = Document(str(DOCX_PATH))

    def apply(s: str) -> str:
        s = s.replace("「", "“").replace("」", "”")
        for old, new in text_rules:
            s = s.replace(old, new)
        return s

    for p in doc.paragraphs:
        if p.text.strip():
            new = apply(p.text)
            if new != p.text:
                p.text = new

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        p.text = apply(p.text)

    doc.save(str(DOCX_PATH))


def main():
    md = MD_PATH.read_text(encoding="utf-8")
    polished = polish_text(md)
    MD_PATH.write_text(polished, encoding="utf-8")
    print(f"Updated MD: {len(polished)} chars")

    rules = [
        ("可答辩、可截图、可叙述创新点的完整形态", "功能完整、便于展示与技术总结的完整形态"),
        ("服务综合实训答辩与课程报告撰写", "支撑综合实训中期检查与后续课程报告撰写"),
        ("答辩环境一键复现", "容器化一键部署"),
        ("适合答辩演示第一印象", "提升首屏视觉与交互体验"),
        ("建议答辩时：用 60% 时间展示本章，40% 时间演示第五章主线。", ""),
        ("（拓展阶段 · 答辩重点）", "（拓展阶段）"),
        ("适合答辩知识关联话题", "便于呈现资源之间的关联关系"),
        ("适合答辩“知识关联”话题", "便于呈现资源之间的关联关系"),
        ("保证答辩现场无 Key 也能演示完整交互", "在未配置 API Key 时仍可完成完整交互演示"),
        ("（便于答辩叙述）", ""),
        ("好演示、好讲解、好写报告", "可用性、可维护性与可扩展性上达到预期目标"),
        ("为期末验收与答辩做好铺垫", "为期末验收与后续功能迭代做好准备"),
        ("答辩 PPT 建议按编号制作功能清单一页", "亦可按功能编号对照《高级功能与报告亮点清单》整理附录"),
        ("后端 REST API 接口清单（主干）", "后端 REST API 接口清单"),
        ("（答辩用）", ""),
        ("报告可写创新点", "创新点说明"),
        ("一键部署答辩环境", "一键部署完整运行环境"),
        ("答辩现场演示脚本建议", "系统演示流程说明"),
        ("答辩常见问题与参考回答（FAQ）", "常见问题说明"),
        ("保证答辩可复现", "保证在无密钥环境下仍可复现"),
        ("演示彩排", "集成验证"),
    ]
    patch_docx(rules)
    print(f"Updated DOCX: {DOCX_PATH}")


if __name__ == "__main__":
    main()
