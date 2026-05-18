# -*- coding: utf-8 -*-
"""生成实训中期报告 Word 文档"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"C:\Users\wanweijie\Desktop\shixun\实训中期报告.docx"


def set_doc_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(12)
    run.bold = bold
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(12)


def main():
    doc = Document()
    set_doc_font(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("软件工程专业综合实训——中期报告")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("软工资源平台（SoftEng Platform）")
    r2.font.size = Pt(14)
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_paragraph()

    # 一、项目背景
    add_heading(doc, "一、项目背景与概述", 1)
    add_para(
        doc,
        "随着软件工程专业课程与实践项目的增多，学院内积累了大量工具链接、课程资料与优秀实训项目，"
        "但资源分散在不同群聊、网盘与个人仓库中，检索困难、版本混乱、难以形成可复用的知识沉淀。"
        "为此，本实训小组设计并实现「软工资源平台」——面向学院师生的资源共享与学习辅助综合平台。"
    )
    add_para(
        doc,
        "平台以「工具资源共享」「课程资源路线」「项目展示与交流」三大业务模块为主线，"
        "配套统一门户首页、用户个人中心与管理员审核能力，"
        "目标是在学院内部建立可搜索、可评价、可审核、可运营的一站式资源入口。"
    )
    add_para(doc, "主要建设目标包括：", bold=True)
    add_bullets(
        doc,
        [
            "工具资源共享与推荐：分类、搜索、收藏与评论，支持成员提交并经管理员审核后上架；",
            "课程资源平台：按学期组织课程，聚合电子资料、网课链接、工具推荐与评论互动；",
            "项目展示与交流：公开展示实训/课程设计项目，支持技术栈标签、GitHub 链接与社区互动；",
            "管理侧能力：待审内容处理、数据统计与后续扩展的智能运营、检索增强等高级能力。",
        ],
    )
    add_para(
        doc,
        "技术架构采用前后端分离：前端 Vue 3 + Vuex + Element Plus；后端 Go + Gin + MySQL；"
        "接口遵循 REST 风格并配套 Swagger 文档；支持 Docker Compose 一键编排与本地 run-dev 脚本快速启动。"
    )

    add_heading(doc, "二、团队分工（简要）", 2)
    add_para(
        doc,
        "万伟杰（组长）：前后端接口设计与对接、代码生成脚本与 Swagger 文档、外链图片本地化；"
        "杨凯宇：主页、登录注册、个人信息与项目模块前端；蔡沛良：工具模块与审核页面；"
        "唐云鹏：课程模块；龚虹宇：管理员审核流、实时统计与默认图标；高颖轩：Service 与 Repository 层及数据访问。"
    )

    # 三、前期已完成
    add_heading(doc, "三、中期前已完成的主要功能", 1)
    add_para(
        doc,
        "截至中期检查前，小组已完成平台核心业务闭环的搭建，用户可在浏览器中完成「浏览—互动—提交—跟踪审核」的主流程。"
        "下列内容为中期前已落地并经过联调验证的能力。"
    )

    add_heading(doc, "3.1 用户与门户", 2)
    add_bullets(
        doc,
        [
            "用户注册、登录、忘记密码；JWT 鉴权；角色区分（学生/教师/管理员）；",
            "交互式门户首页：粒子背景、可折叠侧栏、多搜索引擎、时间显示、常用站点与精选工具/课程/项目入口；",
            "个人中心：资料编辑、收藏管理、提交记录、审核状态、账户安全、系统设置、站内消息（界面与基础交互）。",
        ],
    )

    add_heading(doc, "3.2 工具资源共享模块", 2)
    add_bullets(
        doc,
        [
            "工具列表：分类与标签筛选、按浏览量/收藏量排序、搜索；",
            "工具详情：简介、链接、贡献者、评论/回复/点赞、收藏与浏览量；",
            "工具提交：表单上传，提交后进入待审核；管理员审核通过后方可在列表展示；",
            "外链图片本地化与基于内容哈希的 SVG 默认图标生成。",
        ],
    )

    add_heading(doc, "3.3 课程资源模块", 2)
    add_bullets(
        doc,
        [
            "课程主页：按学期分栏、搜索筛选、课程卡片展示；",
            "课程详情：简介、教师与学分、URL/上传类资源列表、评论互动；",
            "课程资料提交与审核流程对接。",
        ],
    )

    add_heading(doc, "3.4 项目展示模块", 2)
    add_bullets(
        doc,
        [
            "项目列表：分类导航、标签筛选、排序、悬停摘要；",
            "项目详情：Markdown 详情、技术栈、作者、评论、收藏与浏览统计；",
            "项目提交：分步表单、技术栈与标签、GitHub 链接等。",
        ],
    )

    add_heading(doc, "3.5 审核与后端工程化", 2)
    add_bullets(
        doc,
        [
            "审核中心：按工具/课程/项目分模块展示待审条目，支持通过/驳回；",
            "Repository + Service 分层；MySQL 表结构（schema.sql）；",
            "收藏/点赞等计数采用 SQL 子查询实时统计，降低缓存不一致风险；",
            "从 schema 生成 Repository 脚手架、Swagger API 文档集成。",
        ],
    )

    # 四、后期完成（重点）
    add_heading(doc, "四、中期后重点完成与增强功能（本阶段工作）", 1)
    add_para(
        doc,
        "中期之后，小组工作重点从「功能可用」转向「可演示、可答辩、可运维叙事」："
        "在保持三大业务模块稳定的前提下，补齐数据一致性、权限边界、智能检索与运营洞察等能力。"
        "以下为本阶段新增或显著增强的内容（建议作为报告与答辩的重点阐述部分）。",
        bold=True,
    )

    add_heading(doc, "4.1 数据与工程交付", 2)
    add_bullets(
        doc,
        [
            "演示数据集 seed_demo.sql：固定主键的用户、工具、课程、项目、评论、收藏、点赞及待审样本，便于答辩现场展示；",
            "库表补丁 patch_existing_to_demo.sql；Docker Compose 首次启动自动执行 schema + 种子数据；",
            "根目录 run-dev.bat / run-dev.ps1 一键检查库表并启动后端 Air 热重载与前端 dev；cmd/dbcheck 库表自检。",
        ],
    )

    add_heading(doc, "4.2 智能检索、推荐与内容辅助", 2)
    add_bullets(
        doc,
        [
            "统一搜索页：同义词扩展、检索意图识别（自动倾向工具/课程/项目 Tab）；",
            "详情页混合推荐条（热门 + 标签相似 + 共现）；",
            "离线规则摘要与要点（ResourceAiSummary），可扩展为 LLM；",
            "提交前相似标题检测（工具/项目），降低重复投稿；",
            "前后端内容安全校验（敏感词、危险链接协议），覆盖提交与评论写入。",
        ],
    )

    add_heading(doc, "4.3 数据洞察与管理员专属能力", 2)
    add_para(
        doc,
        "下列「洞察」类功能仅对管理员开放：侧栏入口对普通用户隐藏，路由 requiresAdmin 拦截，"
        "与 RAG、WebSocket 演示接口的鉴权策略一致。"
    )
    add_bullets(
        doc,
        [
            "运营数据大屏（ECharts）：提交量、分类分布、趋势等可视化；",
            "资源关联知识图谱：工具—课程—项目关系力导向/环形布局；",
            "前端可观测性面板：HTTP 耗时环缓、Performance API、慢请求样本；",
            "虚拟长列表演示（content-visibility + 万级行 Demo）；",
            "RAG 学习助手：多表检索 + Gemini/OpenAI 生成回答（无 Key 时 demo 模式），引用溯源；",
            "会话画像（本机 localStorage 聚合浏览与搜索，不上传服务器）；",
            "事件时间线：路由浏览、提交、审核等事件合并展示；",
            "审核中心增强：审核步骤与 SLA 示意、表格多选、批量通过/批量驳回；",
            "WebSocket 演示推送（管理员 token 连接）+ 本地通知铃铛。",
        ],
    )

    add_heading(doc, "4.4 体验、国际化与上线向加固", 2)
    add_bullets(
        doc,
        [
            "课程详情学习路径组件：周次甘特条 + 里程碑时间线（演示级教学计划可视化）；",
            "vue-i18n 中英文切换（首页侧栏、分区标题等）；",
            "工具/课程/项目列表：加载中、失败重试、空数据态，避免白屏；",
            "CORS 白名单环境变量、/auth 限流、/live 与 /health 健康检查、Gin release 模式；",
            "生产构建 VUE_APP_API_BASE；注册/改密/改邮等与后端字段路径一致性修复。",
        ],
    )

    add_heading(doc, "4.5 阶段成果小结", 2)
    add_para(
        doc,
        "综上，中期后工作形成了「普通用户：完整资源站 + 搜索与安全提交」与「管理员：审核 + 运营洞察 + RAG」的双层产品结构。"
        "平台已具备本地一键运行、容器化编排、充足演示数据与可截图的高级功能页面，"
        "满足综合实训中期「功能拓展、工程规范与可演示性」的考核要求。"
    )

    add_heading(doc, "五、后续计划（简要）", 1)
    add_bullets(
        doc,
        [
            "完善站内消息、邮箱验证码等真实业务链路（当前部分为演示逻辑）；",
            "课程表 status 字段与待审列表与工具/项目对齐（可选）；",
            "补充自动化测试与 CI；隐私政策等法务静态页（若对外演示需要）。",
        ],
    )

    doc.save(OUT)
    print("已写入:", OUT)


if __name__ == "__main__":
    main()
